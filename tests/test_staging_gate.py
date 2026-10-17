import json

import pytest
from openpyxl import Workbook, load_workbook

from staging_gate import layout as L
from staging_gate.apply import apply_approved
from staging_gate.lock import lock, unprotect, verify
from staging_gate.risks import import_risks
from staging_gate.rollup import rollup

DETECT = "4. DETECT"


def make_wb():
    wb = Workbook()
    wb.active.title = DETECT
    for name in L.LOCKED_SHEETS + L.OPEN_SHEETS:
        if name not in wb.sheetnames:
            wb.create_sheet(name)
    d = wb[DETECT]
    for i, (nr, crit) in enumerate([("D.1", "1-Critical"), ("D.2", "3-Enabler"), ("D.3", "2-Important")]):
        d[f"A{9 + i}"], d[f"H{9 + i}"] = nr, crit
    d["K9"] = "=IFERROR(0,0)"  # category block anchor (rows 9-11)
    cfg = wb[L.CONFIG_SHEET]
    for r, (n, w) in enumerate([("1-Critical", 3), ("2-Important", 2), ("3-Enabler", 1)], 5):
        cfg[f"B{r}"], cfg[f"C{r}"] = n, w
    st = wb[L.STAGING]
    for i, nr in enumerate(["D.1", "D.2", "D.3"]):
        r = 7 + i
        st[f"B{r}"], st[f"C{r}"] = DETECT, nr
        st[f"E{r}"], st[f"F{r}"], st[f"G{r}"] = "state", "proposal", i + 1
        st[f"H{r}"], st[f"I{r}"], st[f"K{r}"] = 2, 3, "ev"
    return wb


def log_rows(wb):
    ws = wb[L.CHANGELOG]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=4, max_col=11) if row[1].value]


def results(wb):
    return [r[10] for r in log_rows(wb)]


def test_only_approved_rows_apply_and_log():
    wb = make_wb()
    wb[L.STAGING]["A7"], wb[L.STAGING]["A9"] = "y", "Y"
    res = apply_approved(wb)
    d = wb[DETECT]
    assert res.controls == 2
    assert (d["S9"].value, d["S10"].value, d["S11"].value) == (1, None, 3)
    assert d["F10"].value is None and d["T9"].value == 2 and d["U9"].value == 3
    assert results(wb).count("APPLIED") == 10
    assert wb[L.STAGING]["O7"].value.endswith(" applied") and wb[L.STAGING]["O8"].value is None
    assert apply_approved(wb).controls == 0  # stamped rows are not reapplied
    assert log_rows(wb)[0][9] == L.STAMP_SOURCE


def test_conflict_partial_stamp_then_override():
    wb = make_wb()
    wb[DETECT]["F9"] = "human text"
    wb[L.STAGING]["A7"] = "y"
    res = apply_approved(wb)
    assert wb[DETECT]["F9"].value == "human text" and res.conflicts == 1
    assert "CONFLICT - skipped" in results(wb)
    assert wb[L.STAGING]["O7"].value.endswith(" partial") and res.controls == 0
    wb[L.STAGING]["A7"] = "y!"
    res = apply_approved(wb)
    assert wb[DETECT]["F9"].value == "state" and res.overrides == 1
    assert any(r[10] == "OVERRIDE" and r[5] == "human text" for r in log_rows(wb))
    assert wb[L.STAGING]["O7"].value.endswith(" overridden")
    assert apply_approved(wb).overrides == 0  # y! does not re-fire


def test_zero_maturity_is_not_a_conflict_but_nonzero_is():
    wb = make_wb()
    wb[DETECT]["S9"], wb[DETECT]["T9"] = 0, 4
    wb[L.STAGING]["A7"] = "y"
    res = apply_approved(wb)
    assert wb[DETECT]["S9"].value == 1 and wb[DETECT]["T9"].value == 4
    assert res.conflicts == 1


def test_out_of_range_maturity_rejected_per_cell():
    wb = make_wb()
    wb[L.STAGING]["A7"], wb[L.STAGING]["G7"] = "y", 9
    res = apply_approved(wb)
    assert wb[DETECT]["S9"].value is None and wb[DETECT]["T9"].value == 2
    assert "REJECTED - not 0-5" in results(wb)
    assert wb[L.STAGING]["O7"].value.endswith(" partial") and res.controls == 0


def test_unknown_sheet_and_nr_logged_not_stamped():
    wb = make_wb()
    st = wb[L.STAGING]
    st["A7"], st["B7"] = "y", "Nope"
    st["A8"], st["C8"] = "y", "ZZ.9"
    res = apply_approved(wb)
    assert res.not_found == 2 and st["O7"].value is None and st["O8"].value is None
    assert set(results(wb)) == {"SHEET NOT FOUND", "NR NOT FOUND"}


def test_lock_and_apply_relocks():
    wb = make_wb()
    locked, missing = lock(wb)
    assert DETECT in locked and not missing
    assert wb[DETECT].protection.sheet and not wb[L.STAGING].protection.sheet
    wb[L.STAGING]["A7"] = "y"
    apply_approved(wb)
    assert wb[DETECT]["S9"].value == 1
    assert all(wb[n].protection.sheet for n in L.LOCKED_SHEETS)
    assert verify(wb[DETECT]) and not verify(wb[DETECT], "wrong")
    with pytest.raises(PermissionError):
        unprotect(wb[DETECT], "wrong")


def test_protection_survives_save_and_wrong_lock_pw_blocks_apply(tmp_path, monkeypatch):
    wb = make_wb()
    lock(wb)
    p = tmp_path / "t.xlsx"
    wb.save(p)
    wb2 = load_workbook(p)
    assert verify(wb2[DETECT])
    wb2[L.STAGING]["A7"] = "y"
    monkeypatch.setattr(L, "LOCK_PW", "different")
    with pytest.raises(PermissionError):
        apply_approved(wb2)
    assert wb2[DETECT]["S9"].value is None


def test_rollup_weighted_and_excludes_zero():
    wb = make_wb()
    wb[L.STAGING]["A7"], wb[L.STAGING]["A9"] = "y", "y"
    apply_approved(wb)
    (b,) = rollup(wb, DETECT)
    assert (b["first"], b["last"]) == (9, 11)
    assert b["rows"] == ["D.1", "D.3"] and b["value"] == pytest.approx((1 * 3 + 3 * 2) / 5)


RISK = {"B": "PR.AA", "C": "Weak MFA", "D": "cur", "E": "tgt", "F": "desc", "G": "act",
        "H": "3-HIGH RISK", "I": "3-MEANINGFUL IMPACT", "J": "1-LOW HANGING FRUIT",
        "M": "CISO", "R": "maturity", "S": "H3", "source_rows": ["D.1", "D.2"],
        "rating_basis": "basis", "rating_evidence": "ev", "confidence": "high"}


def test_risks_import_then_apply(tmp_path):
    wb = make_wb()
    n, rej = import_risks(wb, json.dumps([RISK, {**RISK, "H": "bogus"}, {"B": "x"}]))
    sr = wb[L.STAGING_RISKS]
    assert n == 1 and len(rej) == 2 and "bad H value 'bogus'" in rej[0]
    assert sr["A7"].value is None and sr["Q7"].value is None
    assert sr["N7"].value == "D.1; D.2" and sr["O7"].value == "basis - ev" and sr["P7"].value == "high"
    assert (sr["K7"].value, sr["L7"].value, sr["M7"].value) == ("CISO", "maturity", "H3")
    sr["A7"] = "y"
    assert apply_approved(wb).risks == 1
    r = wb[L.RISKS_SHEET]
    assert r["C12"].value == "Weak MFA" and (r["M12"].value, r["R12"].value, r["S12"].value) == ("CISO", "maturity", "H3")
    assert all(r[f"{c}12"].value is None for c in "AKLN")
    assert sr["Q7"].value.endswith(" applied")
    p = tmp_path / "t.xlsx"
    wb.save(p)
    assert load_workbook(p)[L.RISKS_SHEET]["C12"].value == "Weak MFA"


def test_risks_bad_json_and_single_object():
    with pytest.raises(ValueError, match="did not parse"):
        import_risks(make_wb(), "Here you go: [{}]")
    with pytest.raises(ValueError, match="empty"):
        import_risks(make_wb(), "")
    assert import_risks(make_wb(), json.dumps(RISK))[0] == 1


H3_MD = "# H3\n\n## System prompt (v4.0)\n\n```\nYou are H3.\nJSON only.\n```\n"


def _fix_wb():
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = make_wb()
    ws = wb["Prompts"]
    ws["A7"], ws["B7"], ws["C7"] = "H3 Risk identification", "v1.0", "old prompt"
    dv = DataValidation(type="list", formula1='"y,n"')
    dv.add("A7:A160")
    wb[L.STAGING].add_data_validation(dv)
    lock(wb)
    return wb


def test_fix_workbook_updates_prompt_and_dropdown_and_is_idempotent():
    from staging_gate.fixes import fix_workbook
    wb = _fix_wb()
    res = fix_workbook(wb, H3_MD)
    assert len(res.done) == 2
    assert (wb["Prompts"]["B7"].value, wb["Prompts"]["C7"].value) == ("v4.0", "You are H3.\nJSON only.")
    assert wb[L.STAGING].data_validations.dataValidation[0].formula1 == '"y,n,y!"'
    assert wb["Prompts"].protection.sheet and verify(wb["Prompts"]) and verify(wb[L.CHANGELOG])
    assert results(wb) == ["PROMPT UPDATED", "DROPDOWN UPDATED"]
    again = fix_workbook(wb, H3_MD)
    assert not again.done and len(again.skipped) == 2
    assert len(results(wb)) == 2
