import docx

from staging_gate import profile_review as pr
from staging_gate import review_changes as rc

PROFILE = """# Company profile — ACME TEST OY

> Profile version CP-T-v1, 2026-01-01.

## Part A — Landscape

1. **Identity.** One entity in Helsinki. `[verified: EV-01]`
2. **Scale.** About 260 staff. `[reported: EV-03]`

## Part B — Current state

### DE.CM — Continuous monitoring
Endpoint agent deployed. `[verified: EV-04]` No testing evidence. `[unknown]`

## Part C — Open questions

1. Which log sources reach the SIEM? `[unknown]`

## Part D — Change log

| Date | Change | Trigger | Evidence |
|---|---|---|---|
| 2026-01-01 | Initial | Setup | — |
"""


def build(tmp_path):
    md = tmp_path / "p.md"
    md.write_text(PROFILE, encoding="utf-8")
    assert pr.run({"input": str(md), "stamp": True}) == 0
    out = tmp_path / "review.docx"
    cfg = tmp_path / "c.json"
    cfg.write_text('{"profileVersion": "CP-T-v1", "cisoEmail": "c@x.test", "coordinatorEmail": "p@x.test"}')
    assert pr.run({"input": str(md), "config": str(cfg), "out": str(out), "no_prompt": True}) == 0
    return md, out


def test_ids_stamp_idempotent_and_harvest():
    blocks = pr.assign_ids(pr.parse_markdown(PROFILE))
    ids = [b["claimId"] for b in blocks if b.get("claimId")]
    assert ids == ["A1", "A2", "B.DE.CM.1", "C1"]
    text, added = pr.stamp_markdown(PROFILE, blocks)
    assert added == 4 and "{#A2}" in text and "{#B.DE.CM.1}" in text
    again = pr.assign_ids(pr.parse_markdown(text))
    assert pr.stamp_markdown(text, again)[1] == 0
    meta = pr.harvest(again)
    assert meta["counts"]["unknown"] == 2 and meta["counts"]["verified"] == 2
    assert meta["title"] == "ACME TEST OY"


def test_docx_has_claim_ids_as_first_run(tmp_path):
    _, out = build(tmp_path)
    d = docx.Document(out)
    firsts = [p.runs[0].text.strip() for p in d.paragraphs if p.runs]
    assert {"A1", "A2", "B.DE.CM.1", "C1"} <= set(firsts)


def test_review_roundtrip_extract_apply(tmp_path):
    md, out = build(tmp_path)
    d = docx.Document(out)
    for p in d.paragraphs:
        if p.runs and p.runs[0].text.strip() == "A2":
            d.add_comment(runs=[p.runs[0]], text="Actually 300 staff, see EV-07.", author="R. V", initials="RV")
        if p.runs and p.runs[0].text.strip() == "C1":
            d.add_comment(runs=[p.runs[0]], text="Accept as unknown.", author="R. V", initials="RV")
    d.add_comment(runs=[d.paragraphs[0].runs[0]], text="Cover typo", author="R. V", initials="RV")
    rev = tmp_path / "reviewed.docx"
    d.save(rev)
    props = tmp_path / "props.md"
    assert rc.cmd_extract(str(rev), str(md), str(props)) == 0
    text = props.read_text(encoding="utf-8")
    assert text.count("## CP-") == 3 and "- Claim ID: -" in text and "About 260 staff" in text

    # undecided -> nothing happens
    before = md.read_text(encoding="utf-8")
    assert rc.cmd_apply(str(props), str(md)) == 0 and md.read_text(encoding="utf-8") == before

    marked = []
    for blk in text.split("## CP-")[1:]:
        cid = blk.split("Claim ID: ")[1].splitlines()[0]
        if cid == "A2":
            blk = blk.replace("New text:\n\n\n", "New text:\n\n**Scale.** About 300 staff `[reported: EV-07]`\n\n")
            blk = blk.replace("Approve:", "Approve: y", 1)
        elif cid == "C1":
            blk = blk.replace("Approve:", "Approve: y", 1)
        else:
            blk = blk.replace("Approve:", "Approve: n", 1)
        marked.append(blk)
    props.write_text("## CP-".join([text.split("## CP-")[0]] + marked), encoding="utf-8")
    assert rc.cmd_apply(str(props), str(md), "CP-T-v2") == 0
    after = md.read_text(encoding="utf-8")
    assert "About 300 staff `[reported: EV-07]` {#A2}" in after and "About 260" not in after
    assert "Profile version CP-T-v2" in after
    assert after.count("G1 review — R. V") == 2 and "EV-07" in after.split("Part D")[1]
    assert (tmp_path / "p.md.bak").read_text(encoding="utf-8") == before


def test_unmatched_claim_cannot_be_applied(tmp_path):
    lines = PROFILE.split("\n")
    res = rc.apply_to_profile(PROFILE, [{"id": "CP-001", "claimId": "-", "author": "x", "says": "s",
                                         "newText": "new", "approve": "y"}], None)
    assert res["failed"] and res["text"] == PROFILE and lines
