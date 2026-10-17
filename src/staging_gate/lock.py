"""Sheet protection with Excel's own password scheme (ECMA-376 SHA-512 + salt + spin count),
so a wrong password is refused on unprotect, as Excel's unprotect(LOCK_PW) would."""
import base64
import hashlib
import os
import struct

from openpyxl.utils.protection import hash_password as _legacy_hash

from . import layout as L

SPIN = 100000


def _hash(pw: str, salt: bytes, spin: int) -> str:
    d = hashlib.sha512(salt + pw.encode("utf-16-le")).digest()
    for i in range(spin):
        d = hashlib.sha512(d + struct.pack("<I", i)).digest()
    return base64.b64encode(d).decode()


def verify(ws, pw: str | None = None) -> bool:
    pw = L.LOCK_PW if pw is None else pw
    p = ws.protection
    if p.hashValue:
        if p.algorithmName and p.algorithmName != "SHA-512":
            raise ValueError(f"{ws.title}: unsupported protection hash {p.algorithmName}")
        return _hash(pw, base64.b64decode(p.saltValue), int(p.spinCount or SPIN)) == p.hashValue
    if p.password:
        return _legacy_hash(pw) == p.password
    return True  # protected without a password


def protect(ws, pw: str | None = None):
    pw = L.LOCK_PW if pw is None else pw
    p = ws.protection
    salt = os.urandom(16)
    p._password = None  # openpyxl's setter would hash None; clear the legacy field directly
    p.algorithmName, p.spinCount = "SHA-512", SPIN
    p.saltValue = base64.b64encode(salt).decode()
    p.hashValue = _hash(pw, salt, SPIN)
    p.sheet = True
    p.objects = True
    p.scenarios = True


def unprotect(ws, pw: str | None = None):
    pw = L.LOCK_PW if pw is None else pw
    if not ws.protection.sheet:
        return
    if not verify(ws, pw):
        raise PermissionError(f"{ws.title}: protection password does not match LOCK_PW")
    p = ws.protection
    p.sheet = False
    p._password = None
    p.algorithmName = p.hashValue = p.saltValue = p.spinCount = None


def set_protection(wb, names, on: bool):
    """ApplyApproved.ts setProtection: only touch sheets not already in the wanted state."""
    for name in names:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if on and not ws.protection.sheet:
            protect(ws)
        elif not on and ws.protection.sheet:
            unprotect(ws)


def lock(wb):
    """LockDown: open the staging sheets, (re)protect every live sheet. Returns (locked, missing)."""
    missing = [n for n in L.OPEN_SHEETS + L.LOCKED_SHEETS if n not in wb.sheetnames]
    set_protection(wb, L.OPEN_SHEETS, False)
    locked = []
    for name in L.LOCKED_SHEETS:
        if name in wb.sheetnames:
            ws = wb[name]
            unprotect(ws)
            protect(ws)
            locked.append(name)
    return locked, missing
