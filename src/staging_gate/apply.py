"""ApplyApproved: port of ApplyApproved.ts. Writes only approved staging rows to live sheets."""
from dataclasses import dataclass

from openpyxl.utils import column_index_from_string as ci

from . import layout as L
from .lock import set_protection
from .xl import ChangeLog, cell, first_empty_row, num, s, stamp


@dataclass
class Result:
    controls: int = 0
    overrides: int = 0
    risks: int = 0
    conflicts: int = 0
    not_found: int = 0

    def line(self):
        return (f"Controls applied: {self.controls} | overrides: {self.overrides} | "
                f"risks appended: {self.risks} | conflicts skipped: {self.conflicts} | "
                f"lookup failures: {self.not_found} -- see ChangeLog.")


def _find_nr(ws, nr):
    for r in range(L.FIRST_ROW, L.LAST_ROW + 1):
        v = cell(ws, L.NR_COL, r).value
        if s(v) != "" and s(v) == s(nr):
            return r
    return None


def _blank(v):
    return v is None or v == ""


def _controls(wb, res, log, now):
    st = wb[L.STAGING]
    last = L.STAGING_FIRST_ROW + L.STAGING_ROWS - 1
    for r in range(L.STAGING_FIRST_ROW, last + 1):
        approve = s(cell(st, L.COL_APPROVE, r).value).lower()
        already = s(cell(st, L.COL_STAMP, r).value).lower()
        override, plain = approve == "y!", approve == "y"
        if not (override or plain):
            continue
        if plain and already != "":
            continue
        if override and "overridden" in already:
            continue
        sheet_name, nr = s(cell(st, L.COL_SHEET, r).value), cell(st, L.COL_NR, r).value
        if sheet_name == "" or s(nr) == "":
            continue
        ev = s(cell(st, L.COL_EVIDENCE, r).value)
        conf = s(cell(st, L.COL_CONFIDENCE, r).value)
        if sheet_name not in wb.sheetnames:
            log.write("H2", sheet_name, s(nr), "-", "", "", "", "", "SHEET NOT FOUND")
            res.not_found += 1
            continue
        ws = wb[sheet_name]
        tr = _find_nr(ws, nr)
        if tr is None:
            log.write("H2", sheet_name, s(nr), "-", "", "", "", "", "NR NOT FOUND")
            res.not_found += 1
            continue
        had_conflict = False
        for scol, tcol, is_mat in L.CONTROL_MAP:
            proposed = cell(st, scol, r).value
            if _blank(proposed):
                continue
            target = cell(ws, tcol, tr)
            cur = target.value
            cur_s = s(cur)
            if is_mat:
                clash = (cur_s != "" and cur_s != "0" and num(cur) != 0
                         and num(cur) != num(proposed))
            else:
                clash = cur_s != "" and cur_s != s(proposed)
            if clash and not override:
                log.write("H2", sheet_name, s(nr), tcol, cur_s, s(proposed), ev, conf,
                          "CONFLICT - skipped")
                had_conflict = True
                res.conflicts += 1
                continue
            if is_mat:
                n = num(proposed)
                lo, hi = L.MATURITY_RANGE
                if n != n or n < lo or n > hi:
                    log.write("H2", sheet_name, s(nr), tcol, cur_s, s(proposed), ev, conf,
                              "REJECTED - not 0-5")
                    had_conflict = True
                    continue
                target.value = int(n) if n.is_integer() else n
            else:
                target.value = str(proposed)
            log.write("H2", sheet_name, s(nr), tcol, cur_s, s(proposed), ev, conf,
                      "OVERRIDE" if clash else "APPLIED")
        word = "partial" if had_conflict else ("overridden" if override else "applied")
        cell(st, L.COL_STAMP, r).value = f"{now} {word}"
        if not had_conflict:
            if override:
                res.overrides += 1
            else:
                res.controls += 1


def _risks(wb, res, log, now):
    if L.STAGING_RISKS not in wb.sheetnames or L.RISKS_SHEET not in wb.sheetnames:
        return
    sr, risks = wb[L.STAGING_RISKS], wb[L.RISKS_SHEET]
    nxt = first_empty_row(risks, L.RISK_TARGET_KEY_COL, L.RISK_TARGET_FIRST_ROW)
    last = L.RISK_STAGING_FIRST_ROW + L.RISK_STAGING_ROWS - 1
    for r in range(L.RISK_STAGING_FIRST_ROW, last + 1):
        if s(cell(sr, L.RISK_COL_APPROVE, r).value).lower() != "y":
            continue
        if s(cell(sr, L.RISK_COL_APPLIED, r).value) != "":
            continue
        if s(cell(sr, L.RISK_COL_NAME, r).value) == "":
            continue
        for scol, tcol in L.RISK_APPLY_MAP.items():
            cell(risks, tcol, nxt).value = cell(sr, scol, r).value
        log.write("H3", L.RISKS_SHEET, f"row {nxt}", "B-J,M,R,S", "",
                  s(cell(sr, L.RISK_COL_NAME, r).value), "", s(cell(sr, "P", r).value), "APPLIED")
        cell(sr, L.RISK_COL_APPLIED, r).value = f"{now} applied"
        nxt += 1
        res.risks += 1


def apply_approved(wb) -> Result:
    if L.CHANGELOG not in wb.sheetnames:
        raise ValueError("ChangeLog sheet missing - aborting.")
    res, now = Result(), stamp()
    try:
        set_protection(wb, L.LOCKED_SHEETS, False)
        log = ChangeLog(wb, now)
        if L.STAGING in wb.sheetnames:
            _controls(wb, res, log, now)
        _risks(wb, res, log, now)
    finally:
        set_protection(wb, L.LOCKED_SHEETS, True)
    return res
