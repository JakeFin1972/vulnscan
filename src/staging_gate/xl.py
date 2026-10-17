"""Small openpyxl helpers shared by the commands."""
from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as ci

from . import layout as L


def now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def cell(ws, col: str, row: int):
    return ws.cell(row=row, column=ci(col))


def log(wb, sheet, nr, addr, status, old="", new="", src=""):
    ws = wb[L.CHANGELOG] if L.CHANGELOG in wb.sheetnames else wb.create_sheet(L.CHANGELOG)
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        ws.append(["Timestamp", "Sheet", "nr", "Cell", "Status", "Old", "New", "Staging row"])
    ws.append([now(), sheet, nr, addr, status, old, new, src])


def open_wb(path):
    return load_workbook(path)
