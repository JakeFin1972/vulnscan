"""Port of profile_to_docx.js: render the Company Profile markdown to a Word review copy, with
stable claim IDs (A4, B.DE.CM.2, C1) and an optional --stamp that freezes them into the markdown.
The markdown remains the master; the .docx is generated."""
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

CONTENT_WIDTH = 9000  # DXA: A4 portrait minus 1" margins
TAG_COLOR = {"verified": "2E7D32", "reported": "B26A00", "public-source": "6A6A6A",
             "assumed": "6A6A6A", "unknown": "C62828"}
TAG_RE = r"\[(verified|reported|public-source|assumed|unknown)(:\s*[^\]]*)?\]"
TAG_RE_G = re.compile(TAG_RE, re.I)
ANCHOR_RE = re.compile(r"\s*\{#([A-Za-z0-9._-]+)\}\s*$")
_RULE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_ITEM = re.compile(r"^(?:\d+\.|[-*+])\s+(.*)$")
_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


# ---------------------------------------------------------------- markdown -> blocks
def split_table_row(line: str) -> list[str]:
    return [c.strip() for c in re.sub(r"\|$", "", re.sub(r"^\|", "", line)).split("|")]


def parse_markdown(md: str) -> list[dict]:
    lines = re.split(r"\r?\n", md)
    blocks, i = [], 0
    while i < len(lines):
        t = lines[i].strip()
        if t == "" or _RULE.match(t):
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", t)
        if h:
            blocks.append({"type": "heading", "level": len(h[1]), "text": h[2].strip(), "srcLine": i})
            i += 1
            continue
        if t.startswith("|") and i + 1 < len(lines) and _SEP.match(lines[i + 1]):
            header = split_table_row(t)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i].strip()))
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue
        if t.startswith(">"):
            parts = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                parts.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "note", "text": " ".join(parts).strip()})
            continue
        li = _ITEM.match(t)
        if li:
            text, start = li[1], i
            i += 1
            while i < len(lines):
                n = lines[i].strip()
                if n == "" or _ITEM.match(n) or re.match(r"^#{1,6}\s", n) or n.startswith("|") or _RULE.match(n):
                    break
                text += " " + n
                i += 1
            blocks.append({"type": "item", "text": text.strip(), "srcLine": start})
            continue
        text, start = t, i
        i += 1
        while i < len(lines):
            n = lines[i].strip()
            if (n == "" or re.match(r"^#{1,6}\s", n) or n.startswith("|") or n.startswith(">")
                    or _ITEM.match(n) or _RULE.match(n)):
                break
            text += " " + n
            i += 1
        blocks.append({"type": "para", "text": text.strip(), "srcLine": start})
    return blocks


# ---------------------------------------------------------------- claim IDs
def assign_ids(blocks: list[dict]) -> list[dict]:
    part = category = None
    counters: dict[str, int] = {}
    taken = {m[1] for b in blocks if b.get("text") for m in [ANCHOR_RE.search(b["text"])] if m}

    def next_id(prefix):
        n = counters.get(prefix, 0)
        while True:
            n += 1
            cid = f"{prefix}{n}"
            if cid not in taken:
                break
        counters[prefix] = n
        taken.add(cid)
        return cid

    for b in blocks:
        if b["type"] == "heading":
            if b["level"] <= 2:
                pm = re.match(r"^Part\s+([A-Z])\b", b["text"], re.I)
                part = pm[1].upper() if pm else None
                category = None
                if part and part > "C":
                    part = None
            elif b["level"] == 3 and part == "B":
                cm = re.match(r"^([A-Z]{2}(?:\.[A-Z]{2})?)\b", b["text"])
                category = cm[1] if cm else b["text"].split()[0]
            continue
        if b["type"] in ("table", "note"):
            continue
        is_claim = ((part == "A" and b["type"] == "item") or (part == "C" and b["type"] == "item")
                    or (part == "B" and b["type"] in ("para", "item")))
        if not is_claim:
            continue
        ex = ANCHOR_RE.search(b["text"])
        if ex:
            b["claimId"], b["hadAnchor"] = ex[1], True
            b["text"] = ANCHOR_RE.sub("", b["text"], count=1)
        else:
            prefix = f"B.{category or 'GEN'}." if part == "B" else part
            b["claimId"], b["hadAnchor"] = next_id(prefix), False
        b["part"], b["category"] = part, category
        b["sectionLabel"] = f"Part B / {category or '?'}" if part == "B" else f"Part {part}"
    return blocks


def stamp_markdown(md: str, blocks: list[dict]) -> tuple[str, int]:
    lines = re.split(r"\r?\n", md)
    added = 0
    for b in blocks:
        if not b.get("claimId") or b.get("hadAnchor") or not isinstance(b.get("srcLine"), int):
            continue
        last = b["srcLine"]
        for j in range(b["srcLine"] + 1, len(lines)):
            ln = lines[j]
            if (ln.strip() == "" or re.match(r"^\s*(?:\d+\.|[-*+])\s+", ln) or re.match(r"^\s*#{1,6}\s", ln)
                    or re.match(r"^\s*\|", ln) or re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", ln)):
                break
            last = j
        if ANCHOR_RE.search(lines[last]):
            continue
        lines[last] = lines[last].rstrip() + " {#" + b["claimId"] + "}"
        added += 1
    return "\n".join(lines), added


# ---------------------------------------------------------------- inline formatting
def unwrap_tag_backticks(s: str) -> str:
    return re.sub(r"`(\[[^\]]+\])`", r"\1", s)


def plain_text(text: str) -> str:
    t = unwrap_tag_backticks(text)
    t = re.sub(r"\*\*([^*]+)\*\*", r"\1", t)
    t = re.sub(r"\*([^*]+)\*", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    return ANCHOR_RE.sub("", t).strip()


_INLINE = re.compile(r"\*\*([^*]+)\*\*|(?<!\*)\*([^*]+)\*(?!\*)|`([^`]+)`|" + TAG_RE, re.I)


def _style(run, size=None, bold=None, italics=None, color=None, font=None):
    if size:
        run.font.size = Pt(size / 2)
    if bold is not None:
        run.bold = bold
    if italics is not None:
        run.italic = italics
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if font:
        run.font.name = font


def add_inline(par, text: str, **base):
    rest = unwrap_tag_backticks(text)
    last = 0
    for m in _INLINE.finditer(rest):
        if m.start() > last:
            _style(par.add_run(rest[last:m.start()]), **base)
        if m[1] is not None:
            _style(par.add_run(m[1]), **{**base, "bold": True})
        elif m[2] is not None:
            _style(par.add_run(m[2]), **{**base, "italics": True})
        elif m[3] is not None:
            _style(par.add_run(m[3]), **{**base, "font": "Consolas"})
        else:
            kind = m[4].lower()
            sz = base.get("size") or 22
            _style(par.add_run(" " + m[0] + " "), **{**base, "color": TAG_COLOR.get(kind, "6A6A6A"),
                                                     "bold": kind == "unknown",
                                                     "size": sz if kind == "unknown" else sz - 2})
        last = m.end()
    if last < len(rest):
        _style(par.add_run(rest[last:]), **base)
    if not par.runs:
        _style(par.add_run(rest), **base)
    return par


# ---------------------------------------------------------------- docx helpers
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def _table(doc, ncols, widths):
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t._widths = widths
    return t


def _row(t, values, widths, header=False, shade=None, runs=True):
    cells = t.add_row().cells
    for idx, v in enumerate(values):
        c = cells[idx]
        c.width = Twips(widths[idx])
        _cell_margins(c)
        if header:
            _shade(c, shade or "E8EAED")
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if header:
            _style(p.add_run(plain_text(v)), size=19, bold=True)
        else:
            add_inline(p, str(v), size=19)
    if header:
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))


def build_table(doc, header, rows, widths, header_shade=None):
    t = _table(doc, len(header), widths)
    _row(t, header, widths, header=True, shade=header_shade)
    for r in rows:
        _row(t, r, widths)
    return t


def even_widths(n, total=CONTENT_WIDTH):
    each = total // n
    w = [each] * n
    w[-1] = total - each * (n - 1)
    return w


def callout(doc, fill):
    t = _table(doc, 1, [CONTENT_WIDTH])
    c = t.add_row().cells[0]
    c.width = Twips(CONTENT_WIDTH)
    _shade(c, fill)
    _cell_margins(c)
    return c


def spacer(doc, after=0, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after / 20)
    p.paragraph_format.space_before = Pt(before / 20)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading(doc, text, level):
    p = doc.add_heading(plain_text(text), level=level)
    return p


def _field(par, instr, size=16, color="9A9A9A"):
    run = par.add_run()
    _style(run, size=size, color=color)
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = text
        run._r.append(e)


def contact_line(name, email, missing):
    n, e = (name or "").strip(), (email or "").strip()
    if n and e:
        return f"{n} — {e}"
    if e:
        return e
    if n:
        return f"{n} — [{missing}]"
    return f"[{missing}]"


def _ph(v, fallback):
    return str(v) if v and str(v).strip() else f"[{fallback}]"


def _setup_styles(doc):
    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(11)
    for name, size, color, ital in (("Heading 1", 16, "1A1A1A", False), ("Heading 2", 13, "1A1A1A", False),
                                    ("Heading 3", 11.5, "3C4043", False), ("Heading 4", 11, "3C4043", True)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold, s.font.italic = "Calibri", Pt(size), True, ital
        s.font.color.rgb = RGBColor.from_string(color)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Twips(1440))


# ---------------------------------------------------------------- harvest + build
def harvest(blocks):
    counts = {"total": 0, "verified": 0, "reported": 0, "unknown": 0, "publicSource": 0, "assumed": 0}
    unknowns, sections, title = [], [], "Company Profile"
    for b in blocks:
        if b["type"] == "heading" and b["level"] == 1 and title == "Company Profile":
            title = re.sub(r"^Company profile\s*[—-]\s*", "", plain_text(b["text"]), flags=re.I).strip() or "Company Profile"
        if b["type"] == "heading" and b["level"] == 2:
            t = plain_text(b["text"])
            if re.match(r"^Part\s+[ABC]\b", t, re.I):
                sections.append(t)
        if b["type"] == "heading" and b["level"] == 3:
            t = plain_text(b["text"])
            if re.match(r"^[A-Z]{2}\.[A-Z]{2}\b", t):
                sections.append(t)
        if not b.get("claimId"):
            continue
        counts["total"] += 1
        tags = [m.group(0).lower() for m in TAG_RE_G.finditer(unwrap_tag_backticks(b["text"]))]
        for key, pre in (("verified", "[verified"), ("reported", "[reported"),
                         ("publicSource", "[public-source"), ("assumed", "[assumed")):
            if any(t.startswith(pre) for t in tags):
                counts[key] += 1
        if any(t.startswith("[unknown") for t in tags):
            counts["unknown"] += 1
            unknowns.append({"id": b["claimId"], "section": b.get("sectionLabel", ""), "text": plain_text(b["text"])})
    return {"title": title, "counts": counts, "unknowns": unknowns, "sections": sections}


def build_document(blocks, cfg, meta):
    doc = Document()
    _setup_styles(doc)
    doc.core_properties.author = "H1 Company Profile Creator"
    doc.core_properties.title = meta["title"]
    doc.core_properties.comments = "Company Profile — G1 review copy"
    ciso = contact_line(cfg.get("cisoName"), cfg.get("cisoEmail"), "CISO email not supplied")
    coord = contact_line(cfg.get("coordinatorName"), cfg.get("coordinatorEmail"), "coordinator email not supplied")

    # 1. cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    _style(p.add_run("COMPANY PROFILE"), size=26, bold=True, color="6A6A6A")
    p = doc.add_paragraph()
    _style(p.add_run(meta["title"]), size=44, bold=True)
    c = callout(doc, "FDECEA")
    _style(c.paragraphs[0].add_run("FOR REVIEW — NOT YET APPROVED"), size=24, bold=True, color="C62828")
    _style(c.add_paragraph().add_run(
        "This document is generated from the master profile file. Your comments are proposals against that "
        "master; they are merged into it after the CISO approves them. Editing this Word file does not "
        "change the profile."), size=20)
    spacer(doc, before=300)
    k = meta["counts"]
    build_table(doc, ["", ""], [
        ["Profile version", _ph(cfg.get("profileVersion"), "version not set in config")],
        ["Generated", meta["generated"]],
        ["CISO (closes gate G1)", ciso],
        ["Project coordinator", coord],
        ["Comments due", _ph(cfg.get("reviewDeadline"), "deadline not set in config")],
        ["Gate", "G1 — Profile approval (CISO + one SME per domain)"],
        ["Claims in this document", f"{k['total']}  (verified {k['verified']} · reported {k['reported']} · unknown {k['unknown']})"],
    ], [2600, CONTENT_WIDTH - 2600])
    page_break(doc)

    # 2. how to review
    heading(doc, "How to review this document", 1)
    add_inline(doc.add_paragraph(),
               "Every statement below carries a **claim ID** (A4, B.DE.CM.2, C1) and an **evidence tag**. The ID is "
               "what makes your comment traceable: quote it and the comment can always be matched back to the exact "
               "statement, even after the text around it has changed.")
    c = callout(doc, "F1F3F4")
    _style(c.paragraphs[0].add_run("What we need from you"), size=22, bold=True)
    for txt in (
        "1.  Use **Word comments**, not tracked changes. Start each comment with the claim ID.",
        "2.  For every claim tagged **[unknown]** (red), either supply the answer *and the document that evidences "
        "it*, or state that it is accepted as unknown for now. The gate cannot close until each one has one of "
        "those two outcomes.",
        "3.  Challenge anything tagged **[reported]** (amber). It rests on an interview, not a document. If a "
        "document exists, name it — that is what moves a maturity score.",
        "4.  Comment in " + _ph(cfg.get("returnTo"), "the shared copy — location to be confirmed") +
        ". Please do not circulate separate copies: one commented file is merged reliably, six are not.",
        "Questions about the process or the deadline: **" + coord + "**.  Questions about scope, or to close the "
        "gate: **" + ciso + "**."):
        add_inline(c.add_paragraph(), txt, size=20)
    heading(doc, "Evidence tags", 2)
    build_table(doc, ["Tag", "Means", "Effect on scoring"], [
        ["[verified: EV-nn]", "A named document supports this", "Can support a maturity score"],
        ["[reported: EV-nn]", "Asserted in interview, no document seen", "Caps the score — challenge it"],
        ["[public-source: …]", "From outside the company", "Context only"],
        ["[assumed]", "An explicit assumption", "Context only"],
        ["[unknown]", "Not established", "Blocks scoring — must be closed at G1"]],
        [2200, 3800, CONTENT_WIDTH - 6000])
    heading(doc, "Who reviews what", 2)
    assignments = cfg.get("assignments") or []
    if assignments:
        build_table(doc, ["Section", "Reviewer", "Role"],
                    [[_ph(a.get("scope"), "scope"), _ph(a.get("reviewer"), "unassigned"), _ph(a.get("role"), "role")]
                     for a in assignments], [3000, 3000, CONTENT_WIDTH - 6000])
    else:
        add_inline(doc.add_paragraph(),
                   '*No reviewer assignments were supplied in the config file. Add an "assignments" array to '
                   'review.json so each section has a named owner — an unassigned section is the one that comes '
                   'back unreviewed.*')
    page_break(doc)

    # 3. unknowns register
    heading(doc, "Unknowns register — the G1 checklist", 1)
    add_inline(doc.add_paragraph(),
               "Gate G1 closes when every row below is either answered with evidence or explicitly accepted as "
               "unknown. Nothing else in this document blocks the gate.")
    if meta["unknowns"]:
        build_table(doc, ["ID", "Claim", "Answer + evidence, or “accept as unknown”", "By whom"],
                    [[u["id"], u["text"] if len(u["text"]) <= 240 else u["text"][:237] + "…", "", ""]
                     for u in meta["unknowns"]], [1100, 3900, 2600, CONTENT_WIDTH - 7600])
    else:
        add_inline(doc.add_paragraph(),
                   "*No claims are tagged [unknown]. Either the profile is unusually complete, or the tags are "
                   "missing — worth confirming which before signing off.*")
    page_break(doc)

    # 4. the profile
    heading(doc, "The profile", 1)
    for b in blocks:
        if b["type"] == "heading":
            if b["level"] == 1:
                continue
            heading(doc, b["text"], min(b["level"], 4))
        elif b["type"] == "note":
            add_inline(callout(doc, "FFF8E1").paragraphs[0], b["text"], size=20, italics=True)
            spacer(doc, after=120)
        elif b["type"] == "table":
            build_table(doc, b["header"], b["rows"], even_widths(len(b["header"])))
            spacer(doc, after=160)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            if b.get("claimId"):
                p.paragraph_format.left_indent = Twips(340)
                p.paragraph_format.first_line_indent = Twips(-340)
                # the claim ID must be the FIRST run: review_extract finds claims by it
                _style(p.add_run(b["claimId"] + "  "), size=20, bold=True, color="1A73E8")
            add_inline(p, b["text"])

    # 5. sign-off
    page_break(doc)
    heading(doc, "G1 sign-off", 1)
    add_inline(doc.add_paragraph(),
               "One row per section. “Approved” means: the statements in that section are accurate as far as you "
               "know, and every [unknown] in it has been answered or accepted. Anything less is “Approved with "
               "comments” — list the claim IDs.")
    build_table(doc, ["Section", "Reviewer (name, role)", "Verdict", "Date"],
                [[s, "", "", ""] for s in meta["sections"]], [2600, 3000, 2200, CONTENT_WIDTH - 7800])
    spacer(doc, before=240)
    build_table(doc, ["Gate G1 closed by", "Role", "Profile version approved", "Date"],
                [[ciso, "CISO", _ph(cfg.get("profileVersion"), "version"), ""]], [2600, 2000, 2600, CONTENT_WIDTH - 7200])
    add_inline(doc.add_paragraph(),
               "*Once signed, record the approved version tag in 06_Governance/DecisionLog.md. Only an approved "
               "profile version may be used as input to the maturity assessment.*", size=19, color="6A6A6A")

    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style(hp.add_run(f"{meta['title']}  ·  {_ph(cfg.get('profileVersion'), 'version')}  ·  FOR REVIEW"),
           size=16, color="9A9A9A")
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style(fp.add_run("Page "), size=16, color="9A9A9A")
    _field(fp, "PAGE")
    _style(fp.add_run(" of "), size=16, color="9A9A9A")
    _field(fp, "NUMPAGES")
    return doc


# ---------------------------------------------------------------- command
def _ask_contacts(cfg, args):
    def given(v):
        return v.strip() if v and str(v).strip() else None
    for key, val in (("cisoName", args.get("ciso_name")), ("cisoEmail", args.get("ciso_email")),
                     ("coordinatorName", args.get("coordinator_name")),
                     ("coordinatorEmail", args.get("coordinator_email"))):
        if given(val):
            cfg[key] = val.strip()
    missing = [w for w, k in (("ciso", "cisoEmail"), ("coordinator", "coordinatorEmail")) if not given(cfg.get(k))]
    if not missing:
        return cfg
    if args.get("no_prompt") or not sys.stdin.isatty():
        print(f"Contact email(s) not supplied ({', '.join(missing)}) and prompting is off — "
              "they will appear as [placeholders] in the document.")
        return cfg
    print("\nTwo contacts appear on the review copy so reviewers know who to ask and who signs off.")
    print("Press Enter to leave one blank (it becomes a visible placeholder).\n")
    if "ciso" in missing:
        if not given(cfg.get("cisoName")):
            cfg["cisoName"] = input("CISO name (closes gate G1): ").strip()
        cfg["cisoEmail"] = input("CISO email: ").strip()
    if "coordinator" in missing:
        if not given(cfg.get("coordinatorName")):
            cfg["coordinatorName"] = input("Project coordinator name: ").strip()
        cfg["coordinatorEmail"] = input("Project coordinator email: ").strip()
    return cfg


def run(args: dict) -> int:
    import json
    src = Path(args["input"])
    if not src.exists():
        print(f"Input not found: {src}", file=sys.stderr)
        return 1
    md = src.read_text(encoding="utf-8")
    blocks = assign_ids(parse_markdown(md))
    if args.get("stamp"):
        text, added = stamp_markdown(md, blocks)
        if added == 0:
            print("Every claim already carries an ID anchor — markdown unchanged.")
        else:
            bak = src.with_name(src.name + ".bak")
            bak.write_text(md, encoding="utf-8")
            src.write_text(text, encoding="utf-8")
            print(f"Stamped {added} claim ID(s) into {src.name}  (previous version saved as {bak.name})")
        print("IDs are now frozen: re-running the generator will reuse them even if claims move.")
        return 0
    cfg = {}
    if args.get("config"):
        cp = Path(args["config"])
        if not cp.exists():
            print(f"Config not found: {cp}", file=sys.stderr)
            return 1
        cfg = json.loads(cp.read_text(encoding="utf-8"))
    cfg = _ask_contacts(cfg, args)
    meta = harvest(blocks)
    meta["generated"] = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    out = Path(args.get("out") or re.sub(r"\.md$", "", str(src), flags=re.I) + "_REVIEW.docx")
    build_document(blocks, cfg, meta).save(out)
    k = meta["counts"]
    print(f"Wrote {out}")
    print(f"  claims: {k['total']}  | verified: {k['verified']}  | reported: {k['reported']}  | unknown: {k['unknown']}")
    if k["unknown"]:
        print(f"  unknowns register has {k['unknown']} row(s) — these are the G1 blockers.")
    unstamped = sum(1 for b in blocks if b.get("claimId") and not b.get("hadAnchor"))
    if unstamped:
        print(f"  NOTE: {unstamped} claim(s) have structural (unfrozen) IDs. Run with --stamp once to freeze "
              "them into the markdown before circulating for review.")
    return 0
