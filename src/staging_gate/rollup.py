"""Criticality-weighted average of maturity; rows still at 0 or blank are excluded."""
from openpyxl.utils import column_index_from_string as ci

from . import layout as L
from .xl import blank


def rollup(ws) -> dict:
    num = den = 0.0
    used, skipped = [], []
    for r in range(L.FIRST_ROW, L.LAST_ROW + 1):
        nr = ws.cell(r, ci(L.NR_COL)).value
        m = ws.cell(r, ci(L.MATURITY_COL)).value
        if blank(nr) or blank(m) or m == 0:
            continue
        w = ws.cell(r, ci(L.CRITICALITY_COL)).value
        if isinstance(w, bool) or not isinstance(w, (int, float)) or w <= 0:
            skipped.append(str(nr))
            continue
        num += m * w
        den += w
        used.append(str(nr))
    return {"value": (num / den) if den else 0.0, "rows": used, "no_weight": skipped}
